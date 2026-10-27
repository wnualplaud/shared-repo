"""The mock rows of mock_dynamo.py as one Excel sheet, one account per row, for testers.

    python3 mock_sheet.py [-o mock_accounts.xlsx] [--institution NAME] [--random N] [--seed S]

Use the same --institution / --random / --seed as the mock_dynamo.py run, so
the sheet shows the rows that were loaded. Reads no AWS; needs openpyxl
(pip install openpyxl) and boto3 (imported by mock_dynamo.py).
"""
import argparse

from openpyxl import Workbook
from openpyxl.styles import Font

from mock_dynamo import add_common_args, build

# request keys, then what decides whether the API returns the row, names,
# account, balance, system; the note last. Names are the API's field names.
COLUMNS = [
    "citizenId", "accountId", "accountSubType",
    "ownerType", "accountOwner", "accountStatus", "visibleInApi",
    "accountName", "accountNameEn",
    "accountType", "institutionName", "openingDate", "homeBranch", "homeBranchEn",
    "lastLedgerBalanceAmount", "lastLedgerBalanceCurrency",
    "lastAvailableBalanceAmount", "lastAvailableBalanceCurrency", "creditLimit",
    "accountKey", "src_stm_id", "load_ts", "data_dt",
    "หมายเหตุ",
]


def sheet_row(row, detail, note):
    merged = {**detail, **row}
    merged["visibleInApi"] = "Y" if (row["ownerType"], row["accountOwner"]) == ("INDIVIDUAL", "SINGLE") else "N"
    merged["หมายเหตุ"] = note
    # every cell as text: Excel turns 13-digit ids into 1.23E+12 and drops
    # digits past 15 (9999999999999999.99)
    return [str(merged[c]) if c in merged else "" for c in COLUMNS]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--out", default="mock_accounts.xlsx")
    add_common_args(ap)
    args = ap.parse_args()

    wb = Workbook()
    ws = wb.active
    ws.title = "accounts"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row, detail, note in build(args.institution, args.random, args.seed):
        ws.append(sheet_row(row, detail, note))
    for col in ws.iter_cols(min_row=2):
        for cell in col:
            cell.number_format = "@"
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(60, max(len(str(c.value or "")) for c in col) + 2)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(args.out)
    print(f"wrote {args.out}: {ws.max_row - 1} accounts, institutionName {args.institution!r}")


if __name__ == "__main__":
    main()
